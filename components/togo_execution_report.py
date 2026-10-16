"""The DGBF's standard budget execution tables for one year, as an xlsx file.

One sheet per table, French headers, amounts in whole XOF, rates as
percentages. Every sheet ends with a total row; the sheets that cover all
lines (all but « Investissements », which holds capital lines) add up to the same
totals as « Natures ».
"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from components import togo_execution as te
from components.xlsx_output import workbook_bytes
from constants import translate_func
from translations import t

TITLE_ROW = 1
UNIT_ROW = 2
HEADER_ROW = 3
AMOUNT_FORMAT = "#,##0"
RATE_FORMAT = "0.0%"
_BOLD = Font(bold=True)
_HEADER_FILL = PatternFill("solid", fgColor="DDE4EE")
_WIDTHS = {"code": 12, "text": 48, "amount": 20, "rate": 14}

STAGE_MEASURES = ["dotation_ae", "dotation_initiale", "dotation_finale", "engagement",
                  "liquidation", "ordonnancer"]


def _measure_columns(lang, measures=STAGE_MEASURES):
    columns = [(t(f"execution.measure.{m}", lang), m, "amount") for m in measures]
    return columns + [(t("execution.measure.rate", lang), "rate", "rate")]


def _final_ordered_columns(lang):
    return _measure_columns(lang, ["dotation_finale", "ordonnancer"])


def _write_sheet(wb, name, year, columns, frame, lang, rate_of=("ordonnancer", "dotation_finale")):
    """Write ``frame`` under a title, the unit and a header row, then a total row.

    ``columns`` lists (header, key, kind), kind being "code", "text",
    "amount", "rate" or "share". The total row sums amounts and shares and
    recomputes the rate from the totals of the ``rate_of`` columns.
    """
    ws = wb.create_sheet(name)
    ws.cell(TITLE_ROW, 1, t("execution.report.title", lang, year=year, table=name)).font = _BOLD
    ws.cell(UNIT_ROW, 1, t("execution.report.unit", lang))
    for j, (header, _, kind) in enumerate(columns, start=1):
        cell = ws.cell(HEADER_ROW, j, header)
        cell.font = _BOLD
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(j)].width = _WIDTHS["rate" if kind == "share" else kind]

    for i, record in enumerate(frame.to_dict("records"), start=HEADER_ROW + 1):
        for j, (_, key, kind) in enumerate(columns, start=1):
            _write_value(ws.cell(i, j), record[key], kind, lang)

    total_row = HEADER_ROW + len(frame) + 1
    label_column = next(j for j, (_, _, kind) in enumerate(columns, start=1) if kind == "text")
    ws.cell(total_row, label_column, t("execution.col.total", lang))
    totals = {key: frame[key].sum() for _, key, kind in columns if kind in ("amount", "share")}
    numerator, denominator = rate_of
    for j, (_, key, kind) in enumerate(columns, start=1):
        if kind in ("amount", "share"):
            _write_value(ws.cell(total_row, j), totals[key], kind, lang)
        elif kind == "rate":
            rate = totals[numerator] / totals[denominator] if totals[denominator] else None
            _write_value(ws.cell(total_row, j), rate, kind, lang)
    for cell in ws[total_row]:
        cell.font = _BOLD
    ws.freeze_panes = ws.cell(HEADER_ROW + 1, 1)
    return ws


def _write_value(cell, value, kind, lang):
    if kind == "amount":
        cell.value = int(round(value))
        cell.number_format = AMOUNT_FORMAT
    elif kind in ("rate", "share"):
        cell.value = None if value is None or pd.isna(value) else float(value)
        cell.number_format = RATE_FORMAT
    elif kind == "code":
        cell.value = value or None
    else:
        cell.value = te.label_or_unspecified(value, lang)


def _natures(wb, df, year, lang):
    columns = [
        (t("execution.col.code", lang), "econ1_code", "code"),
        (t("execution.col.nature", lang), "econ1", "text"),
    ] + _measure_columns(lang)
    _write_sheet(wb, t("execution.sheet.natures", lang), year, columns, te.by_nature(df, year), lang)


def _functions(wb, df, year, lang):
    rows = te.by_function(df, year)
    rows["func"] = [translate_func(f, lang) if f else "" for f in rows["func"]]
    columns = [
        (t("execution.col.code", lang), "func_code", "code"),
        (t("execution.col.function", lang), "func", "text"),
    ] + _measure_columns(lang)
    _write_sheet(wb, t("execution.sheet.functions", lang), year, columns, rows, lang)


def _ministries_by_nature(wb, df, year, lang):
    cells = te.ministry_by_nature(df, year)
    natures = te.by_nature(df, year)
    ministries = te.by_ministry(df, year).sort_values("admin2_code")
    frame = ministries[["admin2_code", "admin2"]].reset_index(drop=True)
    final_label = t("execution.measure.dotation_finale", lang)
    ordered_label = t("execution.measure.ordonnancer", lang)
    columns = [
        (t("execution.col.code", lang), "admin2_code", "code"),
        (t("execution.col.ministry", lang), "admin2", "text"),
    ]
    for code, label in zip(natures["econ1_code"], natures["econ1"]):
        name = te.label_or_unspecified(label, lang)
        nature_cells = cells[cells["econ1_code"] == code].set_index("admin2_code")
        for measure, measure_label in (("dotation_finale", final_label), ("ordonnancer", ordered_label)):
            key = f"{measure}:{code}"
            frame[key] = frame["admin2_code"].map(nature_cells[measure]).fillna(0)
            columns.append((f"{name} — {measure_label}", key, "amount"))
    total_label = t("execution.col.total", lang)
    frame["total_finale"] = ministries["dotation_finale"].to_numpy()
    frame["total_ordonnancer"] = ministries["ordonnancer"].to_numpy()
    frame["rate"] = ministries["rate"].to_numpy()
    columns += [
        (f"{total_label} — {final_label}", "total_finale", "amount"),
        (f"{total_label} — {ordered_label}", "total_ordonnancer", "amount"),
        (t("execution.measure.rate", lang), "rate", "rate"),
    ]
    _write_sheet(wb, t("execution.sheet.ministry_nature", lang), year, columns, frame, lang,
                 rate_of=("total_ordonnancer", "total_finale"))


def _programmes(wb, df, year, lang):
    columns = [
        (t("execution.col.ministry_code", lang), "admin2_code", "code"),
        (t("execution.col.ministry", lang), "admin2", "text"),
        (t("execution.col.programme_code", lang), "prog_code", "code"),
        (t("execution.col.programme", lang), "prog", "text"),
        (t("execution.col.nature_code", lang), "econ1_code", "code"),
        (t("execution.col.nature", lang), "econ1", "text"),
    ] + _final_ordered_columns(lang)
    _write_sheet(wb, t("execution.sheet.programmes", lang), year, columns, te.programmes(df, year), lang)


def _projects(wb, df, year, lang):
    columns = [
        (t("execution.col.code", lang), "project_code", "code"),
        (t("execution.col.project", lang), "project", "text"),
        (t("execution.col.funding_code", lang), "funding_code", "code"),
        (t("execution.col.funding", lang), "funding", "text"),
        (t("execution.col.funder", lang), "funder", "text"),
    ] + _measure_columns(lang)
    _write_sheet(wb, t("execution.sheet.projects", lang), year, columns, te.projects(df, year), lang)


def _tag(wb, df, year, dim, lang):
    columns = [
        (t("execution.col.code", lang), "code", "code"),
        (t("execution.col.category", lang), "label", "text"),
        (t("execution.measure.dotation_finale", lang), "dotation_finale", "amount"),
        (t("execution.col.share_finale", lang), "share_finale", "share"),
        (t("execution.measure.ordonnancer", lang), "ordonnancer", "amount"),
        (t("execution.col.share_ordonnancer", lang), "share_ordonnancer", "share"),
        (t("execution.measure.rate", lang), "rate", "rate"),
    ]
    _write_sheet(wb, t(f"execution.sheet.{dim}", lang), year, columns, te.tag_shares(df, year, dim), lang)


def build_report(df, year, lang="fr"):
    """The year's tables from the normalised execution table, as xlsx bytes."""
    if not (df["year"] == year).any():
        raise ValueError(f"no execution lines for {year}")
    wb = Workbook()
    wb.remove(wb.active)
    _natures(wb, df, year, lang)
    _functions(wb, df, year, lang)
    _ministries_by_nature(wb, df, year, lang)
    _programmes(wb, df, year, lang)
    _projects(wb, df, year, lang)
    for dim in te.TAG_DIMENSIONS:
        _tag(wb, df, year, dim, lang)
    return workbook_bytes(wb)
