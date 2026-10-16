"""The DGBF's execution report tables (« TABLEAUX SIGFIP RAPPORT EXECUTION ») for
one budget year, rebuilt from boost.tgo_execution and boost.tgo_execution_detail.

Each sheet keeps the template's rows, columns, French headers and totals, in the
template's order; the period reads « Exercice <year> ». Amounts are whole XOF
written as values, rates are fractions left blank when the allocation is zero.
A single « ORDONNANCER » column holds payment orders on internal resources, as in
the template; payment orders on external resources have their own columns. The
« Notes » sheet states the classification rules and the table left out.
"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from components.xlsx_output import workbook_bytes

AE, CP, ENG, LIQ = "dotation_ae", "dotation_finale", "engagement", "liquidation"
ORD, ORD_RE, ORD_ALL = "ordonnancer_ri", "ordonnancer_re", "ordonnancer"
AMOUNTS = [AE, CP, ENG, LIQ, ORD, ORD_RE, ORD_ALL]

NATURES = [
    ("1", "CHARGES FINANCIERES DE LA DETTE"),
    ("2", "PERSONNEL"),
    ("3", "BIENS ET SERVICES"),
    ("4", "TRANSFERTS COURANTS"),
    ("5", "DEPENSES EN CAPITAL"),
]
FUNCTIONS = [
    ("01", "Services généraux des administrations publiques"),
    ("02", "Défense"),
    ("03", "Ordre et sécurité publics"),
    ("04", "Affaires Economiques"),
    ("05", "Protection de l'environnement"),
    ("06", "Logement et équipements collectifs"),
    ("07", "Santé"),
    ("08", "Loisirs, Culture et Culte"),
    ("09", "Enseignement"),
    ("10", "Protection sociale"),
]
GENDER_SCORES = [("1", "Score Faible"), ("2", "Score Passable"), ("3", "Score Moyen"),
                 ("4", "Score Significatif"), ("5", "Score Principal")]
CLIMATE_CODES = {"11", "21"}
ENVIRONMENT_CODES = {"31", "41", "51", "61"}
SOCIAL_CODE = "2"
# The section whose social lines CHAP_DEP_SOCIALES lists by chapter
SOCIAL_CHAPTERS_SECTION = "210"
CAPITAL = "5"
EXTERNAL = "RE"
OUTSIDE_PIP = "RI_HORS_PROJETS"
UNSPECIFIED = "Non renseigné"

SHEETS = [
    "(GN & MIN)", "FONC PROJET_RI+RE", "PROG_CHAP_PROJ_PIP", "PROG_PROJET", "GN & DOT_PPS_POP",
    "TAB MIN & PROG & DOT", "CF MIN (PREVISIONS)", "CF MIN (EXECUTE)", "CF-MIN_PREV_EXEC",
    "CHAP_DEP_SOCIALES", "SECTEURS SOCIAUX", "AUTRES SECTEURS", "TOUS LES MINISTERES",
    "CF & CE (NOUV)", "TABLEAU 15", "TABLEAU 16", "TABLEAU 17 (RE)", "TABLEAU 18",
    "HORS_PIP & PIP (INTERNES)", "GENRE PAR SCORE", "NEW TAB DBSG_2024", "BSG (PREV_BSG+EXEC_BSG)",
    "BUDGET_VERT",
]

AMOUNT_FORMAT = "#,##0"
RATE_FORMAT = "0.00%"
_BOLD = Font(bold=True)
_HEADER_FILL = PatternFill("solid", fgColor="DDE4EE")
_THIN = Side(style="thin", color="A0A0A0")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEADER_ALIGN = Alignment(wrap_text=True, horizontal="center", vertical="center")
_LABEL_ALIGN = Alignment(wrap_text=True, vertical="top")


# --- Classification ----------------------------------------------------------

def programme_class(label):
    """« dotation », « pilotage » or « operationnel », from the programme label."""
    text = (label or "").strip().upper()
    if text.startswith("DOTATION"):
        return "dotation"
    if text.startswith("PILOTAGE"):
        return "pilotage"
    return "operationnel"


def _prepare(df, year):
    d = df[df["year"] == year].copy()
    if d.empty:
        raise ValueError(f"no execution lines for {year}")
    unknown = sorted(set(d["econ1_code"]) - {code for code, _ in NATURES})
    if unknown:
        raise ValueError(f"economic natures outside the SIGFIP tables in {year}: {unknown}")
    d["class"] = [programme_class(label) for label in d["prog"]]
    d["capital"] = d["econ1_code"] == CAPITAL
    d["external"] = d["funding_code"] == EXTERNAL
    return d


def _functions(d):
    """The ten functions, then any other function code found in the lines."""
    known = [code for code, _ in FUNCTIONS]
    extra = sorted(set(d["func_code"]) - set(known))
    return FUNCTIONS + [(code, code or UNSPECIFIED) for code in extra]


def _sums(d, keys):
    return d.groupby(keys, sort=True)[AMOUNTS].sum()


def _ministries(d):
    """Ministry code -> label, by code."""
    return d.groupby("admin2_code", sort=True)["admin2"].first()


def _ministry_code(code):
    return int(code) if str(code).isdigit() else code


# --- Cells -------------------------------------------------------------------

def _header(ws, row, col, text, to=None):
    """A header cell; ``to`` (row, col) is the other corner of a merged range."""
    end_row, end_col = to or (row, col)
    if (end_row, end_col) != (row, col):
        ws.merge_cells(start_row=row, start_column=col, end_row=end_row, end_column=end_col)
    for r in range(row, end_row + 1):
        for c in range(col, end_col + 1):
            cell = ws.cell(r, c)
            cell.fill = _HEADER_FILL
            cell.border = _BORDER
            cell.font = _BOLD
            cell.alignment = _HEADER_ALIGN
    ws.cell(row, col).value = text


def _text(ws, row, col, value, bold=False, to=None):
    if to:
        ws.merge_cells(start_row=row, start_column=col, end_row=to[0], end_column=to[1])
    cell = ws.cell(row, col, value)
    cell.border = _BORDER
    cell.alignment = _LABEL_ALIGN
    if bold:
        cell.font = _BOLD
    return cell


def _amount(ws, row, col, value, bold=False):
    cell = ws.cell(row, col)
    cell.border = _BORDER
    if bold:
        cell.font = _BOLD
    if value is not None and not pd.isna(value):
        cell.value = int(round(value))
        cell.number_format = AMOUNT_FORMAT
    return cell


def _rate(ws, row, col, numerator, denominator, bold=False):
    cell = ws.cell(row, col)
    cell.border = _BORDER
    if bold:
        cell.font = _BOLD
    if (numerator is not None and denominator is not None and not pd.isna(numerator)
            and not pd.isna(denominator) and denominator != 0):
        cell.value = float(numerator) / float(denominator)
        cell.number_format = RATE_FORMAT
    return cell


def _values(ws, row, col, spec, source, bold=False):
    """Write ``spec`` from column ``col``: a measure name, None for an empty
    column, or ("rate", numerator, denominator). ``source`` None leaves amounts blank."""
    for j, item in enumerate(spec, start=col):
        if item is None:
            _amount(ws, row, j, None, bold)
        elif isinstance(item, tuple):
            _, numerator, denominator = item
            if source is None:
                _rate(ws, row, j, None, None, bold)
            else:
                _rate(ws, row, j, source[numerator], source[denominator], bold)
        else:
            _amount(ws, row, j, None if source is None else source[item], bold)


def _widths(ws, first, rest=18, columns=None):
    ws.column_dimensions["A"].width = first[0]
    if len(first) > 1:
        ws.column_dimensions["B"].width = first[1]
    last = columns or ws.max_column
    for c in range(len(first) + 1, last + 1):
        ws.column_dimensions[get_column_letter(c)].width = rest


def _period(year):
    return f"Exercice {year}"


def _lookup(table, key):
    return table.loc[key] if key in table.index else None


# --- Sheets ------------------------------------------------------------------

_STAGES = [(CP, "DOTATION"), (ENG, "ENGAGEMENT"), (LIQ, "LIQUIDATION"), (ORD, "ORDONNANCER")]


def _gn_min(wb, d, year):
    ws = wb.create_sheet("(GN & MIN)")
    ws["A1"] = _period(year)
    _header(ws, 2, 1, "MINISTERE", to=(3, 2))
    blocks = NATURES + [(None, "TOTAL")]
    for b, (_, label) in enumerate(blocks):
        col = 3 + 4 * b
        _header(ws, 2, col, label, to=(2, col + 3))
        for k, (_, heading) in enumerate(_STAGES):
            _header(ws, 3, col + k, heading)
    cells = _sums(d, ["admin2_code", "econ1_code"])
    ministries = _sums(d, ["admin2_code"])
    natures = _sums(d, ["econ1_code"])
    spec = [m for m, _ in _STAGES]
    row = 4
    for code, label in _ministries(d).items():
        _text(ws, row, 1, _ministry_code(code))
        _text(ws, row, 2, label)
        for b, (nature, _) in enumerate(blocks):
            source = ministries.loc[code] if nature is None else _lookup(cells, (code, nature))
            _values(ws, row, 3 + 4 * b, spec, source)
        row += 1
    _text(ws, row, 1, "Total général", bold=True)
    _text(ws, row, 2, None)
    for b, (nature, _) in enumerate(blocks):
        source = d[AMOUNTS].sum() if nature is None else _lookup(natures, nature)
        _values(ws, row, 3 + 4 * b, spec, source, bold=True)
    _widths(ws, (8, 48), 16)


def _function_table(ws, d, top, spec, functions, label_col=True):
    """One row per function (code and label, or label only), then « Total général »."""
    by_function = _sums(d, ["func_code"])
    start = 3 if label_col else 2
    row = top
    for code, label in functions:
        if label_col:
            _text(ws, row, 1, code)
            _text(ws, row, 2, label)
        else:
            _text(ws, row, 1, label)
        source = _lookup(by_function, code)
        _values(ws, row, start, spec, source if source is not None else pd.Series(0, index=AMOUNTS))
        row += 1
    _text(ws, row, 1, "Total général", bold=True)
    if label_col:
        _text(ws, row, 2, None)
    _values(ws, row, start, spec, d[AMOUNTS].sum(), bold=True)
    return row


def _fonc_projet(wb, d, year):
    ws = wb.create_sheet("FONC PROJET_RI+RE")
    _header(ws, 1, 1, "FONCTION", to=(3, 2))
    _header(ws, 1, 3, _period(year), to=(2, 7))
    for j, heading in enumerate(["DOTATION", "ENGAGEMENT", "LIQUIDATION", "ORDONNANCER",
                                 "Taux d'Exécution"], start=3):
        _header(ws, 3, j, heading)
    projects = d[d["capital"] & (d["funding_code"] != OUTSIDE_PIP)]
    _function_table(ws, projects, 4, [CP, ENG, LIQ, ORD, ("rate", ORD, CP)], _functions(d))
    _widths(ws, (8, 48))


_PROJECT_SPEC = [CP, ENG, LIQ, ORD, ("rate", ORD, CP)]


def _pip_capital(d):
    return d[d["capital"] & (d["funding_code"] != OUTSIDE_PIP)]


def _prog_chap(wb, d, year):
    ws = wb.create_sheet("PROG_CHAP_PROJ_PIP")
    _header(ws, 1, 1, "PROJETS", to=(2, 1))
    _header(ws, 1, 2, _period(year), to=(1, 6))
    for j, heading in enumerate(["DOTATION_FINALE", "ENGAGEMENT", "LIQUIDATION", "ORDONNANCER",
                                 "TAUX D'EXECUTION"], start=2):
        _header(ws, 2, j, heading)
    capital = _pip_capital(d)
    capital = capital[capital["class"] != "dotation"]
    row = 3
    for cls, title in (("pilotage", "PROGRAMMES PILOTAGE ET SOUTIEN"),
                       ("operationnel", "PROGRAMMES OPERATIONNELS")):
        part = capital[capital["class"] == cls]
        _text(ws, row, 1, title, bold=True)
        _values(ws, row, 2, _PROJECT_SPEC, part[AMOUNTS].sum(), bold=True)
        row += 1
        keys = ["admin2_code", "prog_code", "project_code"]
        labels = part.groupby(keys, sort=True)["project"].first()
        for key, source in _sums(part, keys).iterrows():
            _text(ws, row, 1, labels[key] or UNSPECIFIED)
            _values(ws, row, 2, _PROJECT_SPEC, source)
            row += 1
    _text(ws, row, 1, "TOTAL GENERAL", bold=True)
    _values(ws, row, 2, _PROJECT_SPEC, capital[AMOUNTS].sum(), bold=True)
    _widths(ws, (80,))


def _prog_projet(wb, d, year):
    ws = wb.create_sheet("PROG_PROJET")
    capital = _pip_capital(d)
    spec = [CP, ENG, LIQ, ORD]
    sections = [
        ("pilotage", "PROGRAMME (PILOTAGE ET SOUTIEN) PROJETS", 1, 1),
        ("operationnel", "PROGRAMME (OPERATIONNEL) PROJETS", 1, 2),
        ("dotation", "DOTATIONS", 0, 1),
    ]
    row = 1
    for cls, title, gap, after in sections:
        part = capital[capital["class"] == cls]
        _text(ws, row, 1, title, bold=True).border = Border()
        row += 1 + gap
        for j, heading in enumerate(["LIBELLE_PROG", "DOTATION_FINALE", "ENGAGEMENT", "LIQUIDATION",
                                     "ORDONNANCER"], start=1):
            _header(ws, row, j, heading)
        row += 1
        keys = ["admin2_code", "prog_code"]
        labels = part.groupby(keys, sort=True)["prog"].first()
        for key, source in _sums(part, keys).iterrows():
            _text(ws, row, 1, labels[key] or UNSPECIFIED)
            _values(ws, row, 2, spec, source)
            row += 1
        _text(ws, row, 1, "Total général", bold=True)
        _values(ws, row, 2, spec, part[AMOUNTS].sum(), bold=True)
        row += 1 + after
    _header(ws, row, 1, "PROGRAMME", to=(row + 1, 1))
    _header(ws, row, 2, _period(year), to=(row, 6))
    for j, heading in enumerate(["DOTATION_FINALE", "ENGAGEMENT", "LIQUIDATION", "ORDONNANCE",
                                 "TAUX EXECUTION"], start=2):
        _header(ws, row + 1, j, heading)
    row += 2
    programmes = capital[capital["class"] != "dotation"]
    for cls, label in (("pilotage", "Programmes pilotage et soutien"),
                       ("operationnel", "Programmes opérationnels")):
        _text(ws, row, 1, label)
        _values(ws, row, 2, _PROJECT_SPEC, programmes[programmes["class"] == cls][AMOUNTS].sum())
        row += 1
    _text(ws, row, 1, "TOTAL", bold=True)
    _values(ws, row, 2, _PROJECT_SPEC, programmes[AMOUNTS].sum(), bold=True)
    _widths(ws, (80,))


_CLASSES = [("dotation", "DOTATIONS", "Dotation"), ("pilotage", "PROGRAMMES PILOTAGES", "Pilotage"),
            ("operationnel", "PROGRAMMES OPERATIONNELS", "Opérationnel")]


def _pps_pop(wb, d, year):
    ws = wb.create_sheet("GN & DOT_PPS_POP")
    _header(ws, 1, 1, "LIBELLES", to=(2, 2))
    for i, (_, title, _) in enumerate(_CLASSES):
        _header(ws, 1, 3 + 2 * i, title, to=(1, 4 + 2 * i))
    _header(ws, 1, 9, "TOTAL", to=(1, 11))
    for i in range(4):
        _header(ws, 2, 3 + 2 * i, "DOTATION_FINALE")
        _header(ws, 2, 4 + 2 * i, "ORDONNANCER")
    _header(ws, 2, 11, "TAUX EXEC")
    cells = _sums(d, ["econ1_code", "class"])
    natures = _sums(d, ["econ1_code"])
    classes = _sums(d, ["class"])
    labels = dict(NATURES, **{CAPITAL: "DEPENSES EN CAPITAL (RI)"})
    row = 3
    for code, _ in NATURES:
        _text(ws, row, 1, code)
        _text(ws, row, 2, labels[code])
        for i, (cls, _, _) in enumerate(_CLASSES):
            _values(ws, row, 3 + 2 * i, [CP, ORD], _lookup(cells, (code, cls)))
        _values(ws, row, 9, [CP, ORD, ("rate", ORD, CP)], _lookup(natures, code))
        row += 1
    _text(ws, row, 1, "Total général", bold=True, to=(row, 2))
    for i, (cls, _, _) in enumerate(_CLASSES):
        _values(ws, row, 3 + 2 * i, [CP, ORD], _lookup(classes, cls), bold=True)
    _values(ws, row, 9, [CP, ORD, ("rate", ORD, CP)], d[AMOUNTS].sum(), bold=True)

    # Working area under the table: each programme type by nature, then the total
    row += 6
    for cls, _, label in _CLASSES:
        part = d[d["class"] == cls]
        ws.cell(row, 5, label)
        for code, source in _sums(part, ["econ1_code"]).iterrows():
            _values(ws, row, 6, [CP, ORD], source)
            row += 1
        _values(ws, row, 6, [CP, ORD], part[AMOUNTS].sum(), bold=True)
        row += 3
    ws.cell(row, 5, "TOTAL")
    _values(ws, row, 6, [CP, ORD], d[AMOUNTS].sum(), bold=True)
    _widths(ws, (6, 40))


def _tab_min_prog(wb, d, year):
    ws = wb.create_sheet("TAB MIN & PROG & DOT")
    ws["A1"] = _period(year)
    headings = ["MINISTERES / INSTITUTIONS", None, "PROGRAMMES", "DOTATION_AE", "DOTATION_CP", "ENGAGER",
                "LIQUIDER", "ORDONNANCER_RI", "ORDONNANCER_RE", "ORDONNANCER RI+RE"]
    for j, heading in enumerate(headings, start=1):
        _header(ws, 3, j, heading)
    spec = [AE, CP, ENG, LIQ, ORD, ORD_RE, ORD_ALL]
    programmes = _sums(d, ["admin2_code", "prog_code"])
    prog_labels = d.groupby(["admin2_code", "prog_code"], sort=True)["prog"].first()
    ministries = _sums(d, ["admin2_code"])
    row = 4
    for code, label in _ministries(d).items():
        first = True
        for (ministry, prog), source in programmes.loc[[code]].iterrows():
            _text(ws, row, 1, _ministry_code(code) if first else None)
            _text(ws, row, 2, label if first else None)
            _text(ws, row, 3, prog_labels[(ministry, prog)] or UNSPECIFIED)
            _values(ws, row, 4, spec, source)
            first = False
            row += 1
        _text(ws, row, 1, None)
        _text(ws, row, 2, f"Total {label}", bold=True)
        _text(ws, row, 3, None)
        _values(ws, row, 4, spec, ministries.loc[code], bold=True)
        row += 1
    _text(ws, row, 1, None)
    _text(ws, row, 2, "Total général", bold=True)
    _text(ws, row, 3, None)
    _values(ws, row, 4, spec, d[AMOUNTS].sum(), bold=True)
    _widths(ws, (8, 40, 60), 17)


def _cf_min(wb, d, year, name, title, measure):
    ws = wb.create_sheet(name)
    ws["A1"] = title
    ws["C1"] = _period(year)
    functions = _functions(d)
    last = 3 + len(functions)
    _header(ws, 2, 1, "MINISTERES", to=(4, 2))
    _header(ws, 2, 3, "FONCTION", to=(3, last - 1))
    _header(ws, 2, last, "Total", to=(4, last))
    for j, (_, label) in enumerate(functions, start=3):
        _header(ws, 4, j, label)
    cells = _sums(d, ["admin2_code", "func_code"])
    ministries = _sums(d, ["admin2_code"])
    by_function = _sums(d, ["func_code"])
    row = 5
    for code, label in _ministries(d).items():
        _text(ws, row, 1, _ministry_code(code))
        _text(ws, row, 2, label)
        for j, (func, _) in enumerate(functions, start=3):
            _values(ws, row, j, [measure], _lookup(cells, (code, func)))
        _values(ws, row, last, [measure], ministries.loc[code])
        row += 1
    _text(ws, row, 1, "Total général", bold=True)
    _text(ws, row, 2, None)
    for j, (func, _) in enumerate(functions, start=3):
        _values(ws, row, j, [measure], _lookup(by_function, func), bold=True)
    _values(ws, row, last, [measure], d[AMOUNTS].sum(), bold=True)
    _widths(ws, (8, 48), 18)


def _cf_min_prev_exec(wb, d, year):
    ws = wb.create_sheet("CF-MIN_PREV_EXEC")
    functions = _functions(d)
    spec = [AE, CP, ENG, LIQ, ORD]
    headings = ["DOTATION_AE", "DOTATION_FINALE", "ENGAGEMENT_V", "LIQUIDATION", "ORDONNANCE"]
    _header(ws, 1, 1, "MINISTERES", to=(2, 2))
    blocks = functions + [(None, "TOTAL")]
    for b, (_, label) in enumerate(blocks):
        col = 3 + 5 * b
        _header(ws, 1, col, label, to=(1, col + 4))
        for k, heading in enumerate(headings):
            last_block = b == len(blocks) - 1
            _header(ws, 2, col + k, "ORDONNANCER" if last_block and k == 4 else heading)
    cells = _sums(d, ["admin2_code", "func_code"])
    ministries = _sums(d, ["admin2_code"])
    by_function = _sums(d, ["func_code"])
    row = 3
    for code, label in _ministries(d).items():
        _text(ws, row, 1, _ministry_code(code))
        _text(ws, row, 2, label)
        for b, (func, _) in enumerate(blocks):
            source = ministries.loc[code] if func is None else _lookup(cells, (code, func))
            _values(ws, row, 3 + 5 * b, spec, source)
        row += 1
    _text(ws, row, 1, "Total général", bold=True)
    _text(ws, row, 2, None)
    for b, (func, _) in enumerate(blocks):
        source = d[AMOUNTS].sum() if func is None else _lookup(by_function, func)
        _values(ws, row, 3 + 5 * b, spec, source, bold=True)
    _widths(ws, (8, 48), 16)


def _chap_dep_sociales(wb, d_detail):
    ws = wb.create_sheet("CHAP_DEP_SOCIALES")
    for j, heading in enumerate([None, "DOTATION_AE", "DOTATION_CP", "AUTORISATION", "ENGAGEMENT",
                                 "LIQUIDATION", "ORDONNANCE"], start=1):
        _header(ws, 1, j, heading)
    social = d_detail[(d_detail["social_code"] == SOCIAL_CODE)
                      & (d_detail["admin2_code"] == SOCIAL_CHAPTERS_SECTION)]
    spec = [AE, CP, None, ENG, LIQ, ORD]
    labels = social.groupby("admin4_code", sort=True)["admin4"].first()
    row = 2
    for code, source in _sums(social, ["admin4_code"]).iterrows():
        _text(ws, row, 1, labels[code] or UNSPECIFIED)
        _values(ws, row, 2, spec, source)
        row += 1
    _text(ws, row, 1, "TOTAL", bold=True)
    _values(ws, row, 2, spec, social[AMOUNTS].sum(), bold=True)
    _widths(ws, (80,))


_SECTOR_SPEC = [AE, CP, ENG, LIQ, ORD, ("rate", ORD, CP)]


def _sector_header(ws, year, execution_title, period_merged):
    if period_merged:
        _text(ws, 1, 1, _period(year), to=(1, 2)).border = Border()
    else:
        ws["A1"] = _period(year)
    _header(ws, 1, 3, "PREVISIONS", to=(1, 4))
    _header(ws, 1, 5, execution_title, to=(1, 8))
    _header(ws, 2, 1, "MINISTERE", to=(2, 2))
    for j, heading in enumerate(["DOTATION_AE", "DOTATION CP", "ENGAGEMENT", "LIQUIDATION", "ORDONNANCE",
                                 "Taux d'Exécution"], start=3):
        _header(ws, 2, j, heading)


def _ministry_rows(ws, row, d, codes=None):
    labels = _ministries(d)
    sums = _sums(d, ["admin2_code"])
    for code in (codes if codes is not None else labels.index):
        _text(ws, row, 1, _ministry_code(code))
        _text(ws, row, 2, labels[code])
        _values(ws, row, 3, _SECTOR_SPEC, sums.loc[code])
        row += 1
    return row


def _total_row(ws, row, label, source, merge=True):
    _text(ws, row, 1, label, bold=True, to=(row, 2) if merge else None)
    _values(ws, row, 3, _SECTOR_SPEC, source, bold=True)


def _sectors(wb, d, year):
    social = d[d["social_code"] == SOCIAL_CODE]
    other = d[d["social_code"] != SOCIAL_CODE]

    ws = wb.create_sheet("SECTEURS SOCIAUX")
    _sector_header(ws, year, "EXECUTION", period_merged=False)
    row = _ministry_rows(ws, 3, social)
    _total_row(ws, row, "Total général", social[AMOUNTS].sum())
    _widths(ws, (8, 60))

    ws = wb.create_sheet("AUTRES SECTEURS")
    _sector_header(ws, year, "EXECUTIONS", period_merged=True)
    listed = sorted(set(other["admin2_code"]))
    full = d[d["admin2_code"].isin(listed)]
    row = _ministry_rows(ws, 3, full)
    _total_row(ws, row, "Total 1", full[AMOUNTS].sum())
    row += 1
    social_listed = social[social["admin2_code"].isin(listed)]
    for code, source in _sums(social_listed, ["admin2_code"]).iterrows():
        _total_row(ws, row, f"Total 2 (SOCIAUX {code})", source)
        row += 1
    _total_row(ws, row, "TOTAL AUTRES SECTEURS", full[AMOUNTS].sum() - social_listed[AMOUNTS].sum())
    _widths(ws, (8, 60))

    ws = wb.create_sheet("TOUS LES MINISTERES")
    _sector_header(ws, year, "EXECUTIONS", period_merged=True)
    row = _ministry_rows(ws, 3, d)
    _total_row(ws, row, "Total général", d[AMOUNTS].sum())
    _widths(ws, (8, 60))


def _cf_ce(wb, d, year):
    ws = wb.create_sheet("CF & CE (NOUV)")
    functions = _functions(d)
    blocks = functions + [(None, "Total ")]
    rate_col = 2 + 2 * len(blocks)
    _header(ws, 1, 1, "GRANDE_NATURE", to=(2, 1))
    for b, (_, label) in enumerate(blocks):
        col = 2 + 2 * b
        _header(ws, 1, col, label, to=(1, col + 1))
        _header(ws, 2, col, "DOTATION_FINALE")
        _header(ws, 2, col + 1, "ORDONNANCER")
    _header(ws, 1, rate_col, "TAUX NATURE DEPENSES", to=(2, rate_col))
    cells = _sums(d, ["econ1_code", "func_code"])
    natures = _sums(d, ["econ1_code"])
    by_function = _sums(d, ["func_code"])
    row = 3
    for code, label in NATURES:
        _text(ws, row, 1, label)
        for b, (func, _) in enumerate(blocks):
            source = _lookup(natures, code) if func is None else _lookup(cells, (code, func))
            _values(ws, row, 2 + 2 * b, [CP, ORD], source)
        source = _lookup(natures, code)
        _values(ws, row, rate_col, [("rate", ORD, CP)], source)
        row += 1
    total = d[AMOUNTS].sum()
    _text(ws, row, 1, "Total général", bold=True)
    for b, (func, _) in enumerate(blocks):
        source = total if func is None else _lookup(by_function, func)
        _values(ws, row, 2 + 2 * b, [CP, ORD], source, bold=True)
    _values(ws, row, rate_col, [("rate", ORD, CP)], total, bold=True)
    row += 1
    _text(ws, row, 1, "TAUX PAR FONCTIONS", bold=True)
    for b, (func, _) in enumerate(blocks):
        col = 2 + 2 * b
        source = total if func is None else _lookup(by_function, func)
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 1)
        _values(ws, row, col, [("rate", ORD, CP)], source, bold=True)
    _widths(ws, (36,), 16)


def _tableau_15(wb, d, year):
    ws = wb.create_sheet("TABLEAU 15")
    _header(ws, 1, 1, None, to=(2, 1))
    _header(ws, 1, 2, "PREVISIONS", to=(1, 4))
    _header(ws, 1, 5, f"EXECUTIONS {_period(year)}", to=(1, 9))
    for j, heading in enumerate(["DOTATIONS_AE", "DOTATIONS_CP (1)", "AUTORISATION", "ENGAGEMENT (2)",
                                 "TAUX D'EXECUTION (3=2/1)", "LIQUIDATION", "MANDATEMENT (4) ORDONNANCER",
                                 "TAUX D'EXECUTION (5=4/1)"], start=2):
        _header(ws, 2, j, heading)
    spec = [AE, CP, None, ENG, ("rate", ENG, CP), LIQ, ORD, ("rate", ORD, CP)]
    total = d[AMOUNTS].sum()
    natures = _sums(d, ["econ1_code"])
    _text(ws, 3, 1, "DEPENSES BUDGETAIRES", bold=True)
    _values(ws, 3, 2, spec, total, bold=True)
    row = 4
    for code, label in NATURES:
        _text(ws, row, 1, label)
        source = _lookup(natures, code)
        _values(ws, row, 2, spec, source if source is not None else pd.Series(0, index=AMOUNTS))
        row += 1
    _text(ws, row, 1, "TOTAL", bold=True)
    _values(ws, row, 2, spec, total, bold=True)
    _widths(ws, (36,), 18)


_CAPITAL_HEADINGS = ["DOTATION_AE", "DOTATION_CP", "AUTORISATION", "ENGAGEMENT", "LIQUIDATION", "ORDONNANCE",
                     "TAUX D'EXECUTION"]


def _tableau_16(wb, d, year):
    ws = wb.create_sheet("TABLEAU 16")
    _header(ws, 1, 1, "FONCTIONS", to=(2, 1))
    _header(ws, 1, 2, "PREVISIONS (LFI)", to=(1, 4))
    _header(ws, 1, 5, f"EXECUTIONS {_period(year)}", to=(1, 8))
    for j, heading in enumerate(_CAPITAL_HEADINGS, start=2):
        _header(ws, 2, j, heading)
    capital = d[d["capital"]]
    _function_table(ws, capital, 3, [AE, CP, None, ENG, LIQ, ORD, ("rate", ORD, CP)], _functions(d),
                    label_col=False)
    _widths(ws, (48,))


def _tableau_re_ri(wb, d, year, name, title, external):
    ws = wb.create_sheet(name)
    ws["A1"] = title
    _header(ws, 2, 1, "FONCTIONS", to=(3, 1))
    _header(ws, 2, 2, "PREVISIONS (LFI)", to=(2, 4))
    _header(ws, 2, 5, f"EXECUTION {_period(year)}", to=(2, 8))
    measure = ORD_RE if external else ORD
    headings = ["DOTATION_AE", "DOTATION CP", "AUTORISATION", "ENGAGEMENT", "LIQUIDATION",
                "ORDONNANCER_RE" if external else "ORDONNANCER", "TAUX D'EXECUTION"]
    for j, heading in enumerate(headings, start=2):
        _header(ws, 3, j, heading)
    capital = d[d["capital"] & (d["external"] == external)]
    if external:
        present = set(capital["func_code"])
        functions = [(code, label) for code, label in _functions(d) if code in present]
    else:
        functions = _functions(d)
    _function_table(ws, capital, 4, [AE, CP, None, ENG, LIQ, measure, ("rate", measure, CP)], functions,
                    label_col=False)
    _widths(ws, (48,))


def _internal_capital(wb, d, year):
    ws = wb.create_sheet("HORS_PIP & PIP (INTERNES)")
    internal = d[d["capital"] & ~d["external"]]
    blocks = [("INTERNE (HORS PIP)", internal[internal["funding_code"] == OUTSIDE_PIP]),
              ("INTERNE (PIP)", internal[internal["funding_code"] != OUTSIDE_PIP])]
    row = 1
    for title, part in blocks:
        _header(ws, row, 1, title, to=(row + 2, 8))
        _header(ws, row + 3, 1, "FONCTIONS", to=(row + 4, 1))
        _header(ws, row + 3, 2, "PREVISIONS (LFI/LFR)", to=(row + 3, 4))
        _header(ws, row + 3, 5, f"EXECUTION {_period(year)}", to=(row + 3, 8))
        for j, heading in enumerate(["DOTATION_AE", "PREVISIONS", "AUTORISATION", "ENGAGEMENT", "LIQUIDATION",
                                     "ORDONNANCE", "TAUX D'EXECUTION"], start=2):
            _header(ws, row + 4, j, heading)
        total_row = _function_table(ws, part, row + 5, [AE, CP, None, ENG, LIQ, ORD, ("rate", ORD, CP)],
                                    _functions(d), label_col=False)
        row = total_row + 3
    _widths(ws, (48,))


def _scored(d):
    return d[d["gender_code"].isin([code for code, _ in GENDER_SCORES])]


def _gender_by_score(wb, d, year):
    ws = wb.create_sheet("GENRE PAR SCORE")
    ws["A1"] = _period(year)
    blocks = GENDER_SCORES + [(None, "TOTAL SCORES (1+2+3+4+5)")]
    _header(ws, 4, 1, "MINISTERE", to=(4, 2))
    for b, (_, label) in enumerate(blocks):
        col = 3 + 2 * b
        _header(ws, 3, col, label, to=(3, col + 1))
        _header(ws, 4, col, "PREVISIONS")
        _header(ws, 4, col + 1, "EXECUTIONS")
    scored = _scored(d)
    cells = _sums(scored, ["admin2_code", "gender_code"])
    ministries = _sums(scored, ["admin2_code"])
    scores = _sums(scored, ["gender_code"])
    row = 5
    for code, label in _ministries(scored).items():
        _text(ws, row, 1, _ministry_code(code))
        _text(ws, row, 2, label)
        for b, (score, _) in enumerate(blocks):
            source = ministries.loc[code] if score is None else _lookup(cells, (code, score))
            _values(ws, row, 3 + 2 * b, [CP, ORD], source)
        row += 1
    _text(ws, row, 1, "Total général", bold=True)
    _text(ws, row, 2, None)
    for b, (score, _) in enumerate(blocks):
        source = scored[AMOUNTS].sum() if score is None else _lookup(scores, score)
        _values(ws, row, 3 + 2 * b, [CP, ORD], source, bold=True)
    _widths(ws, (8, 48), 16)


def _gender_weights(wb, d, year):
    scored = _scored(d)
    everything = _sums(d, ["admin2_code"])
    genre = _sums(scored, ["admin2_code"]).reindex(everything.index, fill_value=0)
    labels = _ministries(d)

    ws = wb.create_sheet("NEW TAB DBSG_2024")
    _header(ws, 1, 1, "SECTION", to=(2, 1))
    _header(ws, 1, 2, "MINISTERES PILOTES", to=(2, 2))
    _header(ws, 1, 3, f"PREVISIONS {year}", to=(1, 5))
    _header(ws, 1, 6, f"EXECUTION {year} GENRE\nSCORES 1+2+3+4+5\n(4)", to=(2, 6))
    _header(ws, 1, 7, "POIDS \nEXECUTION\nGENRE\n(5)=(4)/(2)", to=(2, 7))
    _header(ws, 2, 3, " BUDGET GLOBAL\nLFI/LFR \n(1)")
    _header(ws, 2, 4, "BUDGET GENRE\nLFI/LFR \n(2)")
    _header(ws, 2, 5, "POIDS\n(3)=(2)/(1)")
    row = 3
    for code in everything.index:
        _text(ws, row, 1, _ministry_code(code))
        _text(ws, row, 2, labels[code])
        _write_weights(ws, row, everything.loc[code], genre.loc[code])
        row += 1
    _text(ws, row, 1, "Total général", bold=True)
    _text(ws, row, 2, None)
    _write_weights(ws, row, d[AMOUNTS].sum(), scored[AMOUNTS].sum(), bold=True)
    _widths(ws, (10, 48), 18)

    ws = wb.create_sheet("BSG (PREV_BSG+EXEC_BSG)")
    _header(ws, 1, 1, "SECTION", to=(2, 1))
    _header(ws, 1, 2, "MINISTERES PILOTES", to=(2, 2))
    _header(ws, 1, 3, f"EXERCICE {year}", to=(1, 4))
    _header(ws, 1, 5, "POIDS \nEXECUTION\nGENRE\n(3)=(2)/(1)", to=(2, 5))
    _header(ws, 2, 3, "BUDGET GENRE\nLFI/LFR \n(1)")
    _header(ws, 2, 4, "EXECUTION GENRE\nSCORES1+2+3+ 4+5\n(2)")
    spec = [CP, ORD, ("rate", ORD, CP)]
    row = 3
    for code in everything.index:
        _text(ws, row, 1, _ministry_code(code))
        _text(ws, row, 2, labels[code])
        _values(ws, row, 3, spec, genre.loc[code])
        row += 1
    _text(ws, row, 1, None)
    _text(ws, row, 2, "Total général", bold=True)
    _values(ws, row, 3, spec, scored[AMOUNTS].sum(), bold=True)
    _widths(ws, (10, 48), 18)


def _write_weights(ws, row, everything, genre, bold=False):
    _amount(ws, row, 3, everything[CP], bold)
    _amount(ws, row, 4, genre[CP], bold)
    _rate(ws, row, 5, genre[CP], everything[CP], bold)
    _amount(ws, row, 6, genre[ORD], bold)
    _rate(ws, row, 7, genre[ORD], genre[CP], bold)


def _green_budget(wb, d, year):
    ws = wb.create_sheet("BUDGET_VERT")
    _header(ws, 1, 1, "SECTION", to=(3, 1))
    _header(ws, 1, 2, "MINISTERES PILOTES", to=(3, 2))
    _header(ws, 1, 3, f"EXECUTION BUDGET VERT EXERCICE {year}", to=(1, 11))
    axes = [("AXE CLIMAT (FAVORABLE)", CLIMATE_CODES), ("AXE ENVIRONNEMENT (FAVORABLE)", ENVIRONMENT_CODES),
            ("TOTAL VERT", CLIMATE_CODES | ENVIRONMENT_CODES)]
    for i, (title, _) in enumerate(axes):
        col = 3 + 3 * i
        _header(ws, 2, col, title, to=(2, col + 2))
        for k, heading in enumerate(["LFI/LFR ", "EXECUTION", "TAUX EXEC\n"]):
            _header(ws, 3, col + k, heading)
    green = d[d["green_code"].isin(CLIMATE_CODES | ENVIRONMENT_CODES)]
    sums = [_sums(green[green["green_code"].isin(codes)], ["admin2_code"]) for _, codes in axes]
    spec = [CP, ORD, ("rate", ORD, CP)]
    zero = pd.Series(0, index=AMOUNTS)
    row = 4
    for code, label in _ministries(green).items():
        _text(ws, row, 1, _ministry_code(code))
        _text(ws, row, 2, label)
        for i, table in enumerate(sums):
            source = _lookup(table, code)
            _values(ws, row, 3 + 3 * i, spec, source if source is not None else zero)
        row += 1
    _text(ws, row, 1, None)
    _text(ws, row, 2, "Total général", bold=True)
    for i, (_, codes) in enumerate(axes):
        _values(ws, row, 3 + 3 * i, spec, green[green["green_code"].isin(codes)][AMOUNTS].sum(), bold=True)
    _widths(ws, (10, 48), 16)


def _notes(wb, d, year):
    ws = wb.create_sheet("Notes")
    total = d[AMOUNTS].sum()
    scores = sorted(set(_scored(d)["gender_code"]))
    lines = [
        ("Tableaux d'exécution du budget (modèle « TABLEAUX SIGFIP RAPPORT EXECUTION »)", True),
        (f"Source : tables boost.tgo_execution et boost.tgo_execution_detail (BOOST Togo), exercice {year}, "
         "lignes du budget général et des comptes spéciaux du Trésor.", False),
        (f"Période : les données sont annuelles ; « du 01/01 au 30/09 » devient « {_period(year)} » et le "
         "préfixe TRIM3 des onglets est retiré.", False),
        ("Montants en francs CFA (XOF), écrits en valeurs ; taux = montant exécuté / dotation, vide quand la "
         "dotation est nulle.", False),
        ("ORDONNANCER, ORDONNANCE, MANDATEMENT et EXECUTION(S) : ordonnancements sur ressources internes "
         "(ORDONNANCER_RI), comme dans le modèle. Les ordonnancements sur ressources externes figurent dans "
         "« TAB MIN & PROG & DOT » (ORDONNANCER_RE, ORDONNANCER RI+RE) et « TABLEAU 17 (RE) ». Les dotations "
         "comprennent les ressources externes.", False),
        (f"Ordonnancements {year} : ressources internes {int(round(total[ORD])):,} + ressources externes "
         f"{int(round(total[ORD_RE])):,} = {int(round(total[ORD_ALL])):,} XOF ; dotation finale "
         f"{int(round(total[CP])):,} XOF.".replace(",", " "), False),
        ("Engagement et liquidation des lignes sur ressources externes : nuls dans les données BOOST ; leurs "
         "dépenses n'apparaissent qu'à l'ordonnancement.", False),
        ("AUTORISATION : absente des données, colonne laissée vide.", False),
        ("Type de programme, d'après son libellé : commençant par « DOTATION » = dotation ; sinon commençant "
         "par « PILOTAGE » = programme de pilotage et soutien ; sinon programme opérationnel.", False),
        ("Projets (FONC PROJET, PROG_CHAP_PROJ_PIP, PROG_PROJET) : dépenses en capital du PIP, c'est-à-dire "
         "financées autrement que sur ressources internes hors projets (projets, fonds de contrepartie, autres "
         "ressources internes, ressources externes), par ligne d'investissement (ADMIN4).", False),
        ("HORS_PIP & PIP (INTERNES) : dépenses en capital sur ressources internes, hors projets d'une part, "
         "projets, fonds de contrepartie et autres ressources internes d'autre part.", False),
        ("Dépenses sociales : lignes CODE_DEPENSES_SOCIALES = 2 (DEPENSES SOCIALES). SECTEURS SOCIAUX : ces "
         "lignes par ministère. AUTRES SECTEURS : les ministères ayant des lignes non sociales, à leur montant "
         "total (Total 1), diminué de leurs lignes sociales (Total 2) ; SECTEURS SOCIAUX + AUTRES SECTEURS = "
         "TOUS LES MINISTERES.", False),
        (f"CHAP_DEP_SOCIALES : lignes sociales de la section {SOCIAL_CHAPTERS_SECTION} (dépenses globales du "
         "ministère de l'Économie et des Finances) par chapitre (ADMIN4) ; son total est la ligne "
         f"{SOCIAL_CHAPTERS_SECTION} de SECTEURS SOCIAUX. Les projets d'autres ministères que le modèle y "
         "ajoute sont comptés dans la ligne de leur ministère de SECTEURS SOCIAUX.", False),
        ("Genre : scores CODE_GENRE 1 (faible) à 5 (principal) ; les données de l'exercice portent les scores "
         f"{', '.join(scores) if scores else 'aucun'} ; les colonnes des autres scores restent vides. "
         "Prévisions = dotation finale.", False),
        ("Budget vert : axe climat = CODE_BUDGET_VERT 11, 21 (adaptation, atténuation favorables) ; axe "
         "environnement = 31, 41, 51, 61 (biodiversité, eau, pollution, déchets favorables).", False),
        ("Fonctions : CODE_FUNC1 01 à 10. Ministères et programmes : codes et libellés des données de "
         "l'exercice.", False),
        ("Onglet non produit : TAB 410. Sa liste d'imputations (subventions aux collectivités territoriales et "
         "investissements de décentralisation de la section 410) est une sélection qu'aucune classification "
         "des données ne permet de reproduire.", False),
        ("Écarts au modèle : CF-MIN_PREV_EXEC suit ses en-têtes (5 mesures par fonction), alors que le modèle "
         "range ses valeurs mesure par mesure sous ces en-têtes ; GN & DOT_PPS_POP reprend sous le tableau la "
         "zone de travail par type de programme, sans les annotations libres ; le modèle compte le ministère "
         "760 dans SECTEURS SOCIAUX et dans AUTRES SECTEURS, ici chaque ligne n'est comptée qu'une fois ; "
         "en-têtes « MINESTERE », « LIQUDATION », « opératinnels » corrigés. Le libellé « DEPENSES EN CAPITAL "
         "(RI) » de GN & DOT_PPS_POP est celui du modèle : la dotation y comprend les ressources externes, "
         "l'ordonnancement est sur ressources internes.", False),
    ]
    for i, (text, bold) in enumerate(lines, start=1):
        cell = ws.cell(i, 1, text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if bold:
            cell.font = _BOLD
    ws.column_dimensions["A"].width = 140


def build_sigfip_report(df_fact, df_detail, year):
    """The year's SIGFIP tables from the normalised fact and detail tables, as xlsx bytes."""
    d = _prepare(df_fact, year)
    detail = df_detail[df_detail["year"] == year]
    wb = Workbook()
    wb.remove(wb.active)
    _gn_min(wb, d, year)
    _fonc_projet(wb, d, year)
    _prog_chap(wb, d, year)
    _prog_projet(wb, d, year)
    _pps_pop(wb, d, year)
    _tab_min_prog(wb, d, year)
    _cf_min(wb, d, year, "CF MIN (PREVISIONS)", "PREVISIONS", CP)
    _cf_min(wb, d, year, "CF MIN (EXECUTE)", "ORDONNANCER", ORD)
    _cf_min_prev_exec(wb, d, year)
    _chap_dep_sociales(wb, detail)
    _sectors(wb, d, year)
    _cf_ce(wb, d, year)
    _tableau_15(wb, d, year)
    _tableau_16(wb, d, year)
    _tableau_re_ri(wb, d, year, "TABLEAU 17 (RE)", "RE", external=True)
    _tableau_re_ri(wb, d, year, "TABLEAU 18", "RI", external=False)
    _internal_capital(wb, d, year)
    _gender_by_score(wb, d, year)
    _gender_weights(wb, d, year)
    _green_budget(wb, d, year)
    _notes(wb, d, year)
    return workbook_bytes(wb)
