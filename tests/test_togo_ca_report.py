"""The compte administratif (spending side) of one ministry, written by build_ca_report."""
import io
import unittest
import zipfile

from openpyxl import load_workbook

from components import togo_execution as te
from components.togo_ca_report import build_ca_report
from tests.execution_detail_fixture import YEAR, detail_frame

SHEETS = ["Dep_prev", "Dep_execut° RECAP", "DEP-execution -DETAILS", "DELNO", "Tab croisé"]


def _content(year=YEAR, ministry="210"):
    return build_ca_report(te.normalise_detail(detail_frame()), year, ministry)


def _rows(ws, label, column=1):
    found = [[c.value for c in row] for row in ws.iter_rows() if row[column - 1].value == label]
    if not found:
        raise AssertionError(f"{label!r} not found in column {column} of {ws.title}")
    return found


def _row(ws, label, column=1):
    return _rows(ws, label, column)[0]


def _cells(ws, *refs):
    return [ws[ref].value for ref in refs]


class CaReportTest(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.wb = load_workbook(io.BytesIO(_content()))

    def test_sheets(self):
        self.assertEqual(self.wb.sheetnames, SHEETS)

    def test_forecast_recap_has_the_template_layout(self):
        ws = self.wb["Dep_prev"]
        self.assertEqual(_cells(ws, "A2", "D2", "E2", "E3", "F3", "E4", "F4"), [
            "Imputation budgétaire", "Libellé du chapitre et de la nature ", "Crédits votés : LFI/LFR",
            "AE", "CP", "(A1)", "(A2)"])
        self.assertEqual([ws.cell(r, 4).value for r in range(5, 11)], [
            "Dépenses de personnel", "Dépenses de biens et services", "Dépenses de transferts courants",
            "Dépenses d'investissement hors PIP", "Dépenses d'investissement sur PIP/RI",
            "Dépenses d'investissement sur PIP/RE"])
        self.assertEqual(_cells(ws, "E5", "F5", "E10", "F10", "A11", "E11", "F11"),
                         [None, 1500, 2000, 1200, "Total du MEF", 3200, 6000])
        self.assertEqual(_cells(ws, "D16", "F16", "E21", "F21", "A23", "E23", "F23"),
                         ["Charges financières dette", 5000, 1500, 1500, "Total des dotations", 1500, 6500])
        self.assertEqual(_cells(ws, "A35", "E35", "F35"),
                         ["Cumul des crédits du MEF (Programmes et dotations)", 4700, 12500])

    def test_execution_recap(self):
        ws = self.wb["Dep_execut° RECAP"]
        self.assertEqual(_cells(ws, "G2", "I2", "K2", "M2", "I4", "L4", "N4"), [
            "Crédits engagés ", "Taux d'engagement", "Liquidation", "Ordonnancement ",
            "C1=B1/A1", "E=D/A2", "G=F/A2"])
        self.assertEqual(_cells(ws, "F5", "H5", "K5", "M5"), [1500, 1100, 1100, 1100])
        self.assertAlmostEqual(ws["J5"].value, 1100 / 1500)
        self.assertIsNone(ws["G5"].value)
        self.assertEqual(_cells(ws, "G10", "H10", "K10", "M10"), [0, 0, 0, 700])
        self.assertAlmostEqual(ws["N10"].value, 700 / 1200)
        self.assertEqual(_cells(ws, "E11", "F11", "G11", "H11", "K11", "M11"),
                         [3200, 6000, 750, 4200, 4140, 4830])
        self.assertAlmostEqual(ws["I11"].value, 750 / 3200)
        self.assertEqual(_cells(ws, "H16", "M16", "F35", "H35", "M35"), [5200, 5200, 12500, 9500, 10130])

    def test_details_by_programme_service_nature_and_line(self):
        ws = self.wb["DEP-execution -DETAILS"]
        self.assertEqual(ws["A1"].value, "Section 210 : MINISTERE DE L'ECONOMIE ET DES FINANCES (MEF)")
        self.assertEqual(ws["A5"].value, "Programme 003 : PILOTAGE ET SOUTIEN AUX SERVICES DU MEF")
        self.assertEqual(ws.print_title_rows, "$2:$4")
        service = _row(ws, "1111080041000")
        self.assertEqual(service[3], "CABINET DU MINISTRE")
        self.assertEqual([service[4], service[6], service[8]], [None, None, None])
        self.assertEqual([service[5], service[7], service[10], service[12]], [1900, 1450, 1440, 1430])
        line = _row(ws, "0031 1111080041000")
        self.assertEqual(line[1:6], ["01101100", "661", "Traitements et salaires en espèces", None, 1500])
        self.assertEqual(line[7], 1100)
        self.assertAlmostEqual(line[9], 1100 / 1500)
        external = [r for r in _rows(ws, "0032 1179001132000") if r[1] == "04104101"][0]
        self.assertEqual(external[2:13], ["211", "Frais de recherche", 2000, 1200, 0, 0, 0.0, 0.0, 0,
                                          0.0, 700])
        total = _row(ws, "Total des dépenses du programme 003")
        self.assertEqual([total[4], total[5], total[6], total[7], total[10], total[12]],
                         [2900, 3700, 500, 1950, 1890, 2580])
        self.assertEqual(_row(ws, "Dotation 900 : CHARGES FINANCIERES DE LA DETTE PUBLIQUE")[0],
                         "Dotation 900 : CHARGES FINANCIERES DE LA DETTE PUBLIQUE")
        final = _row(ws, "Cumul des crédits du MEF (Programmes et dotations)")
        self.assertEqual([final[4], final[5], final[7], final[12]], [4700, 12500, 9500, 10130])

    def test_unpaid_commitments(self):
        ws = self.wb["DELNO"]
        self.assertEqual(_cells(ws, "A2", "C2", "E2", "K2", "I5", "J5", "K4", "L4"), [
            "Référence", "Crédits votés LFI/LFR", "Crédits engagés", "Reste à payer budgétaire",
            "E= C-B2", "F= D-C", "G1= D-B1", "G2= D-B2"])
        self.assertEqual(ws["A6"].value, "Prog : 003")
        self.assertEqual(_row(ws, "Total programme 003")[2:12],
                         [2900, 3700, 500, 1950, 1890, 2580, -60, 690, 650, 630])
        self.assertEqual(_row(ws, "Total section 210")[2:12],
                         [3200, 6000, 750, 4200, 4140, 4830, -60, 690, 650, 630])
        self.assertEqual(_row(ws, "Total dotations")[2:12],
                         [1500, 6500, 100, 5300, 5300, 5300, 0, 0, 0, 0])
        external = _row(ws, "Investissement PIP/RE", column=2)
        self.assertEqual(external[2:12], [2000, 1200, 0, 0, 0, 700, 0, 700, 700, 700])

    def test_programme_by_nature_cross_table(self):
        ws = self.wb["Tab croisé"]
        self.assertEqual(_cells(ws, "A4", "B4", "D4", "J4", "D5", "L5", "D6", "E6"), [
            "Numéro programme", "Intitulé des Programmes ", "Dépenses courantes", "Dépenses en capital",
            "Dépenses de personnel", "Dép. d'investissement sur PIP", "Prévision", "Exécution"])
        self.assertEqual(_row(ws, "003")[3:13],
                         [1500, 1100, 400, 330, None, None, None, None, 1800, 1150])
        self.assertEqual(_row(ws, "Total")[3:13], [1500, 1100, 400, 330, 2000, 2000, 300, 250, 1800, 1150])

    def test_metadata_carries_no_generator(self):
        with zipfile.ZipFile(io.BytesIO(_content())) as archive:
            core = archive.read("docProps/core.xml").decode()
            app = archive.read("docProps/app.xml").decode()
        self.assertNotIn("creator", core)
        self.assertNotIn("lastModifiedBy", core)
        self.assertNotIn("openpyxl", app.lower())

    def test_ministry_without_lines_is_refused(self):
        with self.assertRaises(ValueError):
            _content(ministry="999")
        with self.assertRaises(ValueError):
            _content(year=YEAR + 5)


if __name__ == "__main__":
    unittest.main()
