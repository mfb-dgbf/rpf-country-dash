"""scripts/togo_execution_report.py writes one year's tables from the database."""
import importlib.util
import pathlib
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from queries import QueryService
from tests import execution_detail_fixture as lines
from tests.execution_fixture import TOTAL_FINALE, YEAR, execution_frame

SCRIPT = pathlib.Path(__file__).resolve().parent.parent / "scripts" / "togo_execution_report.py"


def _load_script():
    spec = importlib.util.spec_from_file_location("togo_execution_report_cli", SCRIPT)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


class CliTest(unittest.TestCase):

    def tearDown(self):
        QueryService._instance = None

    @patch("queries.PUBLIC_ONLY", False)
    def test_writes_the_year_from_a_fresh_read(self):
        cli = _load_script()
        with tempfile.TemporaryDirectory() as tmp, \
                patch.object(QueryService, "get_togo_execution_data",
                             return_value=execution_frame()) as read:
            out = pathlib.Path(tmp) / "report.xlsx"
            self.assertEqual(cli.main(["--year", str(YEAR), "--out", str(out)]), 0)
            ws = load_workbook(out)["Natures"]
        read.assert_called_once_with(persistent=False)
        self.assertEqual(ws.cell(ws.max_row, 5).value, TOTAL_FINALE)

    @patch("queries.PUBLIC_ONLY", False)
    def test_year_without_lines_fails(self):
        cli = _load_script()
        with tempfile.TemporaryDirectory() as tmp, \
                patch.object(QueryService, "get_togo_execution_data", return_value=execution_frame()):
            out = pathlib.Path(tmp) / "report.xlsx"
            self.assertEqual(cli.main(["--year", "1990", "--out", str(out)]), 1)
            self.assertFalse(out.exists())

    @patch("queries.PUBLIC_ONLY", False)
    def test_writes_the_sigfip_tables(self):
        cli = _load_script()
        with tempfile.TemporaryDirectory() as tmp, \
                patch.object(QueryService, "get_togo_execution_data", return_value=lines.fact_frame()), \
                patch.object(QueryService, "get_togo_execution_detail",
                             return_value=lines.detail_frame()) as read:
            out = pathlib.Path(tmp) / "sigfip.xlsx"
            self.assertEqual(cli.main(["--year", str(lines.YEAR), "--sigfip", "--out", str(out)]), 0)
            wb = load_workbook(out)
        read.assert_called_once_with(lines.YEAR, persistent=False)
        self.assertEqual(wb.sheetnames[0], "(GN & MIN)")
        total = [c.value for c in wb["TOUS LES MINISTERES"][wb["TOUS LES MINISTERES"].max_row]]
        self.assertEqual(total[3], lines.TOTAL_FINALE)

    @patch("queries.PUBLIC_ONLY", False)
    def test_writes_one_ministry_compte_administratif(self):
        cli = _load_script()
        mef = lines.detail_frame()
        mef = mef[mef["admin2_code"] == "210"]
        with tempfile.TemporaryDirectory() as tmp, \
                patch.object(QueryService, "get_togo_execution_data") as fact, \
                patch.object(QueryService, "get_togo_execution_detail", return_value=mef) as read:
            out = pathlib.Path(tmp) / "ca.xlsx"
            self.assertEqual(cli.main(["--year", str(lines.YEAR), "--ca", "--ministry", "210",
                                       "--out", str(out)]), 0)
            wb = load_workbook(out)
        read.assert_called_once_with(lines.YEAR, "210", persistent=False)
        fact.assert_not_called()
        self.assertEqual(wb.sheetnames, ["Dep_prev", "Dep_execut° RECAP", "DEP-execution -DETAILS", "DELNO",
                                         "Tab croisé"])

    @patch("queries.PUBLIC_ONLY", False)
    def test_ca_needs_a_ministry(self):
        cli = _load_script()
        with self.assertRaises(SystemExit):
            cli.main(["--year", "2025", "--ca", "--out", "ca.xlsx"])


if __name__ == "__main__":
    unittest.main()
