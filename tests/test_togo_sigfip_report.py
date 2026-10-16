"""The DGBF's execution report tables (« TABLEAUX SIGFIP ») written by build_sigfip_report."""
import io
import unittest
import zipfile

from openpyxl import load_workbook

from components import togo_execution as te
from components.togo_sigfip_report import SHEETS, build_sigfip_report
from tests.execution_detail_fixture import (
    TOTAL_AE,
    TOTAL_ENGAGEMENT,
    TOTAL_FINALE,
    TOTAL_LIQUIDATION,
    TOTAL_ORDONNANCER,
    TOTAL_ORDONNANCER_RE,
    TOTAL_ORDONNANCER_RI,
    YEAR,
    detail_frame,
    fact_frame,
)

EXPECTED_SHEETS = [
    "(GN & MIN)", "FONC PROJET_RI+RE", "PROG_CHAP_PROJ_PIP", "PROG_PROJET", "GN & DOT_PPS_POP",
    "TAB MIN & PROG & DOT", "CF MIN (PREVISIONS)", "CF MIN (EXECUTE)", "CF-MIN_PREV_EXEC",
    "CHAP_DEP_SOCIALES", "SECTEURS SOCIAUX", "AUTRES SECTEURS", "TOUS LES MINISTERES",
    "CF & CE (NOUV)", "TABLEAU 15", "TABLEAU 16", "TABLEAU 17 (RE)", "TABLEAU 18",
    "HORS_PIP & PIP (INTERNES)", "GENRE PAR SCORE", "NEW TAB DBSG_2024", "BSG (PREV_BSG+EXEC_BSG)",
    "BUDGET_VERT", "Notes",
]


def _content(year=YEAR):
    return build_sigfip_report(te.normalise_execution(fact_frame()),
                               te.normalise_detail(detail_frame()), year)


def _values(ws, row):
    return [c.value for c in ws[row]]


def _row(ws, label, column=1):
    """The values of the first row whose cell in ``column`` is ``label``."""
    for row in ws.iter_rows():
        if row[column - 1].value == label:
            return [c.value for c in row]
    raise AssertionError(f"{label!r} not found in column {column} of {ws.title}")


class SigfipReportTest(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.wb = load_workbook(io.BytesIO(_content()))

    def test_sheets_in_template_order_without_tab_410(self):
        self.assertEqual(self.wb.sheetnames, EXPECTED_SHEETS)
        self.assertEqual(SHEETS, EXPECTED_SHEETS[:-1])

    def test_ministries_by_nature(self):
        ws = self.wb["(GN & MIN)"]
        self.assertEqual(ws["A1"].value, f"Exercice {YEAR}")
        self.assertEqual(ws["A2"].value, "MINISTERE")
        self.assertEqual([ws.cell(2, c).value for c in (3, 7, 11, 15, 19, 23)], [
            "CHARGES FINANCIERES DE LA DETTE", "PERSONNEL", "BIENS ET SERVICES",
            "TRANSFERTS COURANTS", "DEPENSES EN CAPITAL", "TOTAL"])
        self.assertEqual(_values(ws, 3)[2:6], ["DOTATION", "ENGAGEMENT", "LIQUIDATION", "ORDONNANCER"])
        self.assertEqual(_row(ws, 210)[:6], [210, "MINISTERE DE L'ECONOMIE ET DES FINANCES",
                                             5000, 5200, 5200, 5200])
        # no debt line for health: the cells stay blank
        self.assertEqual(_row(ws, 610)[2:6], [None, None, None, None])
        total = _row(ws, "Total général")
        self.assertEqual(total[22:26], [TOTAL_FINALE, TOTAL_ENGAGEMENT, TOTAL_LIQUIDATION,
                                        TOTAL_ORDONNANCER_RI])
        self.assertEqual(total[18:22], [4400, 850, 800, 800])

    def test_ministry_programmes_split_ordered_amounts(self):
        ws = self.wb["TAB MIN & PROG & DOT"]
        self.assertEqual(_values(ws, 3), [
            "MINISTERES / INSTITUTIONS", None, "PROGRAMMES", "DOTATION_AE", "DOTATION_CP", "ENGAGER",
            "LIQUIDER", "ORDONNANCER_RI", "ORDONNANCER_RE", "ORDONNANCER RI+RE"])
        first = _row(ws, 210)
        self.assertEqual(first[1:3], ["MINISTERE DE L'ECONOMIE ET DES FINANCES",
                                      "PILOTAGE ET SOUTIEN AUX SERVICES DU MEF"])
        self.assertEqual(first[3:], [2900, 3700, 1950, 1890, 1880, 700, 2580])
        mef = _row(ws, "Total MINISTERE DE L'ECONOMIE ET DES FINANCES", column=2)
        self.assertEqual(mef[4], 12500)
        total = _row(ws, "Total général", column=2)
        self.assertEqual(total[3:], [TOTAL_AE, TOTAL_FINALE, TOTAL_ENGAGEMENT, TOTAL_LIQUIDATION,
                                     TOTAL_ORDONNANCER_RI, TOTAL_ORDONNANCER_RE, TOTAL_ORDONNANCER])

    def test_capital_by_function_and_funding(self):
        t16 = _row(self.wb["TABLEAU 16"], "Total général")
        t17 = _row(self.wb["TABLEAU 17 (RE)"], "Total général")
        t18 = _row(self.wb["TABLEAU 18"], "Total général")
        self.assertEqual(t16[1:3], [5500, 4400])
        self.assertEqual(t17[1:3], [2800, 2000])
        self.assertEqual(t18[1:3], [2700, 2400])
        self.assertEqual(t17[6], TOTAL_ORDONNANCER_RE)
        self.assertEqual(t16[6], t18[6])
        self.assertIsNone(t16[3])
        self.assertEqual([r[0].value for r in self.wb["TABLEAU 17 (RE)"].iter_rows(min_row=4)],
                         ["Affaires Economiques", "Santé", "Total général"])
        self.assertAlmostEqual(_row(self.wb["TABLEAU 17 (RE)"], "Santé")[7], 200 / 800)

    def test_internal_capital_outside_and_in_the_pip(self):
        ws = self.wb["HORS_PIP & PIP (INTERNES)"]
        self.assertEqual(ws["A1"].value, "INTERNE (HORS PIP)")
        totals = [r for r in ws.iter_rows(values_only=True) if r[0] == "Total général"]
        self.assertEqual([list(r[1:7]) for r in totals], [
            [300, 300, None, 250, 250, 250],
            [2400, 2100, None, 600, 550, 550],
        ])

    def test_natures_table_15(self):
        ws = self.wb["TABLEAU 15"]
        self.assertEqual(_values(ws, 2)[1:], [
            "DOTATIONS_AE", "DOTATIONS_CP (1)", "AUTORISATION", "ENGAGEMENT (2)",
            "TAUX D'EXECUTION (3=2/1)", "LIQUIDATION", "MANDATEMENT (4) ORDONNANCER",
            "TAUX D'EXECUTION (5=4/1)"])
        self.assertEqual(ws["E1"].value, f"EXECUTIONS Exercice {YEAR}")
        budget = _row(ws, "DEPENSES BUDGETAIRES")
        self.assertEqual(budget[1:], _row(ws, "TOTAL")[1:])
        self.assertEqual(budget[2], TOTAL_FINALE)
        self.assertAlmostEqual(budget[8], TOTAL_ORDONNANCER_RI / TOTAL_FINALE)
        self.assertEqual(_row(ws, "DEPENSES EN CAPITAL")[1:3], [5500, 4400])

    def test_projects_by_programme_type(self):
        ws = self.wb["PROG_CHAP_PROJ_PIP"]
        self.assertEqual(_row(ws, "PROGRAMMES PILOTAGE ET SOUTIEN")[1:5], [1800, 500, 450, 450])
        self.assertEqual(_row(ws, "PROJET BONNE GOUVERNANCE FINANCIERE")[1], 1800)
        self.assertEqual(_row(ws, "PROGRAMMES OPERATIONNELS")[1:5], [800, 0, 0, 0])
        self.assertEqual(_row(ws, "TOTAL GENERAL")[1], 2600)
        labels = [r[0] for r in ws.iter_rows(values_only=True)]
        self.assertNotIn("PROG D'APPUI AUX POPULATIONS VULNERABLES", labels)

        ws = self.wb["PROG_PROJET"]
        self.assertEqual(_row(ws, "DOTATION : DEPENSES COMMUNES D'INVESTISSEMENT")[1], 1500)
        self.assertEqual(_row(ws, "TOTAL")[1:5], [2600, 500, 450, 450])

        fonc = _row(self.wb["FONC PROJET_RI+RE"], "Total général")
        self.assertEqual(fonc[2:6], [4100, 600, 550, 550])

    def test_programme_types_by_nature(self):
        ws = self.wb["GN & DOT_PPS_POP"]
        self.assertEqual(_values(ws, 2)[2:], ["DOTATION_FINALE", "ORDONNANCER"] * 4 + ["TAUX EXEC"])
        self.assertEqual(_row(ws, "CHARGES FINANCIERES DE LA DETTE", column=2)[2:10],
                         [5000, 5200, None, None, None, None, 5000, 5200])
        self.assertEqual(_row(ws, "PERSONNEL", column=2)[4:6], [2200, 1750])
        total = _row(ws, "Total général")
        self.assertEqual(total[2:10], [6500, 5300, 4400, 2530, 3190, 2340, TOTAL_FINALE,
                                       TOTAL_ORDONNANCER_RI])

    def test_social_sectors_add_up_to_all_ministries(self):
        social = _row(self.wb["SECTEURS SOCIAUX"], "Total général")
        others = _row(self.wb["AUTRES SECTEURS"], "TOTAL AUTRES SECTEURS")
        everyone = _row(self.wb["TOUS LES MINISTERES"], "Total général")
        self.assertEqual(social[2:7], [3200, 5600, 3250, 3200, 3200])
        self.assertEqual(everyone[2:7], [TOTAL_AE, TOTAL_FINALE, TOTAL_ENGAGEMENT, TOTAL_LIQUIDATION,
                                         TOTAL_ORDONNANCER_RI])
        self.assertEqual([s + o for s, o in zip(social[2:7], others[2:7])], everyone[2:7])
        ws = self.wb["AUTRES SECTEURS"]
        self.assertEqual(_row(ws, "Total 1")[3], TOTAL_FINALE)
        self.assertEqual(_row(ws, "Total 2 (SOCIAUX 210)")[3], 4100)
        self.assertEqual(_row(ws, "Total 2 (SOCIAUX 610)")[3], 1500)

        chapters = self.wb["CHAP_DEP_SOCIALES"]
        self.assertEqual(_values(chapters, 1)[1:], ["DOTATION_AE", "DOTATION_CP", "AUTORISATION",
                                                     "ENGAGEMENT", "LIQUIDATION", "ORDONNANCE"])
        self.assertEqual(_row(chapters, "UNITE DE POLITIQUE FISCALE")[2], 2000)
        self.assertEqual(_row(chapters, "TOTAL")[1:], [2400, 4100, None, 2600, 2550, 2550])

    def test_functions_by_ministry_and_nature(self):
        ws = self.wb["CF MIN (PREVISIONS)"]
        self.assertEqual(ws["C4"].value, "Services généraux des administrations publiques")
        self.assertEqual(ws["L4"].value, "Protection sociale")
        self.assertEqual(_row(ws, 210)[2:13], [9800, None, None, 1200, None, None, None, None,
                                              None, 1500, 12500])
        self.assertEqual(_row(self.wb["CF MIN (EXECUTE)"], "Total général")[12], TOTAL_ORDONNANCER_RI)

        wide = self.wb["CF-MIN_PREV_EXEC"]
        self.assertEqual(_values(wide, 2)[2:7], ["DOTATION_AE", "DOTATION_FINALE", "ENGAGEMENT_V",
                                                 "LIQUIDATION", "ORDONNANCE"])
        health = _row(wide, 610)
        self.assertEqual(health[32:37], [800, 1590, 740, 740, 740])
        self.assertEqual(_row(wide, "Total général")[52:57], [
            TOTAL_AE, TOTAL_FINALE, TOTAL_ENGAGEMENT, TOTAL_LIQUIDATION, TOTAL_ORDONNANCER_RI])

        cross = self.wb["CF & CE (NOUV)"]
        self.assertEqual(_row(cross, "DEPENSES EN CAPITAL")[1:3], [900, 700])
        self.assertEqual(_row(cross, "Total général")[21:23], [TOTAL_FINALE, TOTAL_ORDONNANCER_RI])
        self.assertAlmostEqual(_row(cross, "TAUX PAR FONCTIONS")[21], TOTAL_ORDONNANCER_RI / TOTAL_FINALE)

    def test_gender_and_green_budget(self):
        ws = self.wb["GENRE PAR SCORE"]
        self.assertEqual([ws.cell(3, c).value for c in (3, 5, 7, 9, 11, 13)], [
            "Score Faible", "Score Passable", "Score Moyen", "Score Significatif", "Score Principal",
            "TOTAL SCORES (1+2+3+4+5)"])
        self.assertEqual(_row(ws, 210)[2:], [None] * 6 + [600, 450, 1200, 0, 1800, 450])
        self.assertEqual(_row(ws, "Total général")[12:14], [2600, 450])
        dbsg = _row(self.wb["NEW TAB DBSG_2024"], "Total général")
        self.assertEqual(dbsg[2:4] + dbsg[5:6], [TOTAL_FINALE, 2600, 450])
        bsg = _row(self.wb["BSG (PREV_BSG+EXEC_BSG)"], "Total général", column=2)
        self.assertEqual(bsg[2:4], [2600, 450])
        self.assertEqual(self.wb["NEW TAB DBSG_2024"]["C1"].value, f"PREVISIONS {YEAR}")

        green = self.wb["BUDGET_VERT"]
        self.assertEqual(green["C2"].value, "AXE CLIMAT (FAVORABLE)")
        self.assertEqual(_row(green, 210)[2:11], [600, 450, 0.75, 1200, 0, 0, 1800, 450, 0.25])
        self.assertEqual(_row(green, "Total général", column=2)[8:10], [2600, 450])

    def test_notes_state_the_rules_and_the_left_out_table(self):
        text = "\n".join(str(r[0]) for r in self.wb["Notes"].iter_rows(values_only=True) if r[0])
        for fragment in ("TAB 410", "DOTATION", "PILOTAGE", "CODE_DEPENSES_SOCIALES", "11, 21",
                         "31, 41, 51, 61", "AUTORISATION", "ressources internes"):
            self.assertIn(fragment, text)

    def test_amount_and_rate_formats(self):
        ws = self.wb["TOUS LES MINISTERES"]
        total = [c for c in ws[ws.max_row]]
        self.assertIsInstance(total[3].value, int)
        self.assertEqual(total[3].number_format, "#,##0")
        self.assertEqual(total[7].number_format, "0.00%")

    def test_metadata_carries_no_generator(self):
        with zipfile.ZipFile(io.BytesIO(_content())) as archive:
            core = archive.read("docProps/core.xml").decode()
            app = archive.read("docProps/app.xml").decode()
        self.assertNotIn("creator", core)
        self.assertNotIn("lastModifiedBy", core)
        self.assertNotIn("openpyxl", app.lower())

    def test_year_without_lines_is_refused(self):
        with self.assertRaises(ValueError):
            _content(YEAR + 5)


if __name__ == "__main__":
    unittest.main()
