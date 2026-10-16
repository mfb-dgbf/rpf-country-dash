"""The DGBF standard tables written to Excel by build_report."""
import io
import unittest
import zipfile

from openpyxl import load_workbook

from components import togo_execution as te
from components.togo_execution_report import HEADER_ROW, build_report
from tests.execution_fixture import (
    OPEN_YEAR,
    TOTAL_FINALE,
    TOTAL_INITIALE,
    TOTAL_ORDONNANCER,
    YEAR,
    execution_frame,
)

SHEETS = ["Natures", "Fonctions", "Ministères × natures", "Programmes", "Investissements",
          "Genre", "Budget vert", "Dépenses sociales"]


def _workbook(year=YEAR):
    return load_workbook(io.BytesIO(build_report(te.normalise_execution(execution_frame()), year)))


def _headers(ws):
    return [c.value for c in ws[HEADER_ROW]]


def _column(ws, header):
    return _headers(ws).index(header) + 1


def _body(ws):
    """Data rows (no header, no total row), as lists of values."""
    rows = [[c.value for c in row] for row in ws.iter_rows(min_row=HEADER_ROW + 1)]
    return rows[:-1]


def _total(ws, header):
    return ws.cell(row=ws.max_row, column=_column(ws, header)).value


class ReportTest(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.wb = _workbook()

    def test_one_sheet_per_table(self):
        self.assertEqual(self.wb.sheetnames, SHEETS)

    def test_natures_totals_and_headers(self):
        ws = self.wb["Natures"]
        self.assertEqual(_headers(ws), [
            "Code", "Nature économique", "Dotation AE", "Dotation initiale",
            "Dotation finale", "Engagement", "Liquidation", "Ordonnancement", "Taux d'exécution",
        ])
        self.assertEqual(ws.cell(row=ws.max_row, column=2).value, "Total")
        self.assertEqual(_total(ws, "Dotation finale"), TOTAL_FINALE)
        self.assertEqual(_total(ws, "Dotation initiale"), TOTAL_INITIALE)
        self.assertEqual(_total(ws, "Ordonnancement"), TOTAL_ORDONNANCER)
        self.assertEqual(_total(ws, "Engagement"), 430)
        self.assertAlmostEqual(_total(ws, "Taux d'exécution"), TOTAL_ORDONNANCER / TOTAL_FINALE)

    def test_amounts_are_whole_francs_and_rates_percentages(self):
        ws = self.wb["Natures"]
        amount = ws.cell(row=HEADER_ROW + 1, column=_column(ws, "Dotation finale"))
        rate = ws.cell(row=HEADER_ROW + 1, column=_column(ws, "Taux d'exécution"))
        self.assertIsInstance(amount.value, int)
        self.assertEqual(amount.number_format, "#,##0")
        self.assertEqual(rate.number_format, "0.0%")

    def test_every_sheet_adds_up_to_the_year(self):
        for name in ["Fonctions", "Programmes", "Genre", "Budget vert", "Dépenses sociales"]:
            ws = self.wb[name]
            self.assertEqual(_total(ws, "Dotation finale"), TOTAL_FINALE, name)
            self.assertEqual(_total(ws, "Ordonnancement"), TOTAL_ORDONNANCER, name)

    def test_ministries_by_nature_is_wide_with_totals(self):
        ws = self.wb["Ministères × natures"]
        self.assertIn("PERSONNEL — Dotation finale", _headers(ws))
        self.assertIn("DEPENSES EN CAPITAL — Ordonnancement", _headers(ws))
        self.assertEqual(_total(ws, "Total — Dotation finale"), TOTAL_FINALE)
        self.assertEqual(_total(ws, "Total — Ordonnancement"), TOTAL_ORDONNANCER)
        mef = next(r for r in _body(ws) if r[0] == "210")
        self.assertEqual(mef[_column(ws, "DEPENSES EN CAPITAL — Dotation finale") - 1], 650)

    def test_projects_hold_the_capital_lines(self):
        ws = self.wb["Investissements"]
        self.assertEqual(_total(ws, "Dotation finale"), 650)
        self.assertEqual(_total(ws, "Ordonnancement"), 460)
        funder = next(r for r in _body(ws) if r[0] == "P2")[_column(ws, "Bailleur") - 1]
        self.assertEqual(funder, "BANQUE MONDIALE")

    def test_functions_in_french(self):
        labels = [r[1] for r in _body(self.wb["Fonctions"])]
        self.assertIn("Services publics généraux", labels)
        self.assertIn("Santé", labels)
        self.assertIn("Non renseigné", labels)

    def test_undefined_rate_is_blank(self):
        ws = self.wb["Programmes"]
        capital = next(r for r in _body(ws) if r[2] == "010" and r[4] == "5")
        self.assertIsNone(capital[_column(ws, "Taux d'exécution") - 1])

    def test_metadata_carries_no_generator(self):
        content = build_report(te.normalise_execution(execution_frame()), YEAR)
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            core = archive.read("docProps/core.xml").decode()
            app = archive.read("docProps/app.xml").decode()
        self.assertNotIn("creator", core)
        self.assertNotIn("lastModifiedBy", core)
        self.assertNotIn("openpyxl", app.lower())
        self.assertNotIn("<Application>", app)

    def test_year_without_rows_is_refused(self):
        with self.assertRaises(ValueError):
            _workbook(OPEN_YEAR + 5)


if __name__ == "__main__":
    unittest.main()
